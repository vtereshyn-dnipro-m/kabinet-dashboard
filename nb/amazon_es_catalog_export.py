# Databricks notebook source
# MAGIC %md
# MAGIC # Kabinet - Amazon ES Catalog Export
# MAGIC
# MAGIC Выгрузка для Марко: все наши активные товары на Amazon ES — SKU, ASIN, EAN и ссылки на все
# MAGIC изображения (главное, затем PT01, PT02… отдельными колонками) — в Excel. Запускается
# MAGIC вручную кнопкой «Run now» у джобы; расписания нет.
# MAGIC
# MAGIC **Откуда что берётся**
# MAGIC - Перечень товаров — сегодняшний отчёт листингов Amazon (`raw_amazon_get_merchant_listings_all_data`,
# MAGIC   его грузит джоба Дарины «FBA Fees + Listings» в 06:15), только `Spain (ES)` и статус `Active`.
# MAGIC   Строки `amzn.gr.…` — это SKU, которые Amazon сам заводит под перепродажу возвращённых
# MAGIC   товаров (Grade & Resell), а не наши, поэтому их в выгрузке нет.
# MAGIC - EAN и изображения — **свежий запрос** SP-API Catalog Items 2022-04-01 (`searchCatalogItems`
# MAGIC   по 20 ASIN), рынок ES. Не старые снимки витрины.
# MAGIC - EAN, которого нет в каталоге Amazon, берём из справочника SKU Кабинета (`kabinet_data.sku_master`)
# MAGIC   и пишем отдельной колонкой, откуда он. Код ищем как есть и в очищенном виде (без `-FBA`,
# MAGIC   `-FBM`, хвостового `_`, «` 2`» → «`-2`»), но НЕ откатываем к базовому коду: у варианта `-B`
# MAGIC   и у набора `S2_` свой штрихкод, и чужой EAN хуже пустого.
# MAGIC - Изображения: у каждого варианта Amazon отдаёт несколько размеров — берём самый большой
# MAGIC   (по площади). Порядок колонок: главное (MAIN), затем PT01, PT02… по номеру. Образцы цвета
# MAGIC   (SWCH) и прочие служебные варианты в колонки не идут — их число видно в сводке.
# MAGIC
# MAGIC **Куда** — `/Volumes/dnipro_m/dnipro_m/kabinet_raw/exports/amazon_es_catalog/`: файл с датой
# MAGIC выгрузки и копия `…_latest.xlsx`. Пульса и правила живости у джобы нет намеренно: она не по
# MAGIC расписанию, и правило «нет успеха N часов» горело бы каждый день, когда выгрузку просто не
# MAGIC просили.
# MAGIC
# MAGIC Параметр `limit` — сколько товаров взять (0 = все); нужен для пробы.

# COMMAND ----------

# Ячейка 1: вся работа. Самодостаточна.
#
# Ячейки с `%pip` нет намеренно: зависимости объявлены в спецификации окружения джобы, а
# `restartPython()` рядом с ней роняет прогон SIGABRT (AGENTS.md).
import io, json, re, shutil, time
from datetime import datetime, timezone, timedelta

import psycopg2

# Файл для Марко — на английском (решение владельца 06.10.2026): заголовки, значения «откуда EAN»,
# названия листов и сводка. Значения источника — константы, по ним же считается сводка.
SRC_AMAZON, SRC_DIR, SRC_NONE = "Amazon catalog", "SKU directory", "none"
import requests
from databricks.sdk import WorkspaceClient
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

PG_HOST = "ep-delicate-cherry-d2nabn27.database.us-east-1.cloud.databricks.com"
PG_EP = "projects/kabinet-dashboard/branches/production/endpoints/primary"
EP = "https://sellingpartnerapi-eu.amazon.com"
MP_ES = "A1RKKUPIHCS9HS"
OUT_DIR = "/Volumes/dnipro_m/dnipro_m/kabinet_raw/exports/amazon_es_catalog"
# searchCatalogItems: 2 запроса в секунду с запасом 2. Берём с запасом — 0,6 с между запросами.
PAUSE_SEC = 0.6

try:
    dbutils.widgets.text("limit", "0")  # noqa: F821
    LIMIT = int(dbutils.widgets.get("limit") or 0)  # noqa: F821
except Exception:
    LIMIT = 0

KYIV = timezone(timedelta(hours=3))
now = datetime.now(KYIV)
stamp = now.strftime("%Y-%m-%d")
stamp_human = now.strftime("%Y-%m-%d %H:%M")

# ── 1. активные листинги ES из сегодняшнего отчёта ─────────────────────────────
lst = spark.sql("""
    WITH last AS (SELECT max(report_date) AS d
                  FROM dnipro_m.dnipro_m.raw_amazon_get_merchant_listings_all_data
                  WHERE marketplace_name = 'Spain (ES)')
    SELECT DISTINCT l.seller_sku,
           coalesce(l.asin1, CASE WHEN l.product_id_type = '1' THEN l.product_id END) AS asin,
           l.fulfillment_channel, l.report_date
    FROM dnipro_m.dnipro_m.raw_amazon_get_merchant_listings_all_data l, last
    WHERE l.marketplace_name = 'Spain (ES)' AND l.status = 'Active' AND l.report_date = last.d
      AND l.seller_sku NOT LIKE 'amzn.gr.%'
""").toPandas()  # noqa: F821
listing_date = str(lst["report_date"].iloc[0]) if len(lst) else "—"
lst = lst.sort_values("seller_sku").reset_index(drop=True)
n_grade = spark.sql("""
    SELECT count(*) FROM dnipro_m.dnipro_m.raw_amazon_get_merchant_listings_all_data
    WHERE marketplace_name = 'Spain (ES)' AND status = 'Active' AND seller_sku LIKE 'amzn.gr.%'
      AND report_date = (SELECT max(report_date) FROM dnipro_m.dnipro_m.raw_amazon_get_merchant_listings_all_data
                         WHERE marketplace_name = 'Spain (ES)')
""").collect()[0][0]  # noqa: F821
if LIMIT:
    lst = lst.head(LIMIT)
print(f"листингов ES: {len(lst)} (отчёт от {listing_date}; amzn.gr.* исключено: {n_grade})")
no_asin = lst[lst["asin"].isna()]
asins = sorted(set(lst["asin"].dropna()))

# ── 2. SP-API Catalog Items ────────────────────────────────────────────────────
def sec(k):
    return dbutils.secrets.get("amazon-sp-api", k)  # noqa: F821

_lwa = requests.post("https://api.amazon.com/auth/o2/token", data={
    "grant_type": "refresh_token", "refresh_token": sec("refresh-token"),
    "client_id": sec("lwa-app-id"), "client_secret": sec("lwa-client-secret")}, timeout=30)
_lwa.raise_for_status()
HDR = {"x-amz-access-token": _lwa.json()["access_token"]}

def search(batch):
    """Один запрос на ≤20 ASIN. На 429 ждём нарастающе, до ~5 минут в сумме, потом падаем."""
    params = {"identifiers": ",".join(batch), "identifiersType": "ASIN", "marketplaceIds": MP_ES,
              "includedData": "identifiers,images,summaries", "pageSize": 20}
    wait = 2.0
    for _ in range(9):
        r = requests.get(f"{EP}/catalog/2022-04-01/items", headers=HDR, params=params, timeout=60)
        if r.status_code == 429:
            time.sleep(wait); wait = min(wait * 2, 60); continue
        r.raise_for_status()
        return r.json().get("items", [])
    raise RuntimeError(f"[QUOTA_429] searchCatalogItems: квота не освободилась, батч {batch[:3]}…")

items = {}
for i in range(0, len(asins), 20):
    for it in search(asins[i:i + 20]):
        items[it["asin"]] = it
    time.sleep(PAUSE_SEC)
print(f"ASIN запрошено {len(asins)}, в каталоге найдено {len(items)}")

def ean_from_catalog(it):
    found = []
    for blk in it.get("identifiers", []):
        if blk.get("marketplaceId") != MP_ES:
            continue
        for x in blk.get("identifiers", []):
            t, v = x.get("identifierType"), str(x.get("identifier") or "").strip()
            if t == "EAN" and v:
                found.append(v)
            elif t == "GTIN" and len(v) == 14 and v.startswith("0"):
                found.append(v[1:])          # GTIN-14 с ведущим нулём — тот же EAN-13
    return list(dict.fromkeys(found))

def pt_key(variant):
    if variant == "MAIN":
        return (0, 0)
    m = re.fullmatch(r"PT(\d+)", variant or "")
    return (1, int(m.group(1))) if m else None

def images(it):
    """{вариант: ссылка на самый большой размер}; служебные варианты считаем отдельно."""
    best, other = {}, set()
    for blk in it.get("images", []):
        if blk.get("marketplaceId") != MP_ES:
            continue
        for im in blk.get("images", []):
            v = im.get("variant")
            if pt_key(v) is None:
                other.add(v); continue
            area = (im.get("height") or 0) * (im.get("width") or 0)
            if v not in best or area > best[v][0]:
                best[v] = (area, im.get("link"), im.get("width"), im.get("height"))
    return best, other

def title(it):
    for s in it.get("summaries", []):
        if s.get("marketplaceId") == MP_ES:
            return s.get("itemName")
    return None

# ── 3. справочник SKU Кабинета — запасной EAN ──────────────────────────────────
_w = WorkspaceClient()
conn = psycopg2.connect(host=PG_HOST, port=5432, dbname="databricks_postgres",
                        user=_w.current_user.me().user_name,
                        password=_w.postgres.generate_database_credential(endpoint=PG_EP).token,
                        sslmode="require")
cur = conn.cursor()
cur.execute("SELECT sku, ean FROM kabinet_data.sku_master WHERE ean IS NOT NULL AND ean <> ''")
dir_ean = dict(cur.fetchall())
conn.close()

def candidates(sku):
    """Варианты записи того же кода — без отката к базовому: у варианта и набора свой EAN."""
    s = sku.strip()
    c = [s]
    t = re.sub(r"-(FBA|FBM)_?$", "", s)
    t = re.sub(r"_+$", "", t)
    c += [t, t + "_", re.sub(r" (\d)$", r"-\1", t)]
    out = []
    for x in c:
        out.append(x)
        m = re.fullmatch(r"(S\d*_)?(\d{5,7})(-.*)?", x)   # Amazon местами теряет ведущий ноль
        if m:
            out.append((m.group(1) or "") + m.group(2).zfill(8) + (m.group(3) or ""))
    return list(dict.fromkeys(out))

def ean_from_dir(sku):
    for c in candidates(sku):
        if dir_ean.get(c):
            return dir_ean[c], c
    return None, None

# ── 4. строки выгрузки ─────────────────────────────────────────────────────────
rows, other_variants = [], {}
for r in lst.itertuples():
    it = items.get(r.asin) if isinstance(r.asin, str) else None
    eans = ean_from_catalog(it) if it else []
    src, dir_code = None, None
    if eans:
        ean, src = ", ".join(eans), SRC_AMAZON
    else:
        ean, dir_code = ean_from_dir(r.seller_sku)
        src = SRC_DIR if ean else SRC_NONE
    best, other = images(it) if it else ({}, set())
    for v in other:
        other_variants[v] = other_variants.get(v, 0) + 1
    order = sorted(best, key=pt_key)
    rows.append({
        "sku": r.seller_sku, "asin": r.asin or "", "title": title(it) if it else None,
        "in_catalog": bool(it), "ean": ean or "", "ean_src": src, "ean_multi": len(eans) > 1,
        "imgs": {v: best[v][1] for v in order}, "n_img": len(order),
        "fc": "FBA" if str(r.fulfillment_channel).startswith("AMAZON") else "FBM",
    })

pt_max = max([max([pt_key(v)[1] for v in x["imgs"] if v != "MAIN"] or [0]) for x in rows] or [0])
img_cols = ["MAIN"] + [f"PT{n:02d}" for n in range(1, pt_max + 1)]

# ── 5. Excel ───────────────────────────────────────────────────────────────────
wb = Workbook()
ws = wb.active
ws.title = "Amazon ES"
head = (["SKU", "ASIN", "EAN", "EAN source"]
        + ["Main image" if c == "MAIN" else c for c in img_cols] + ["Export date"])
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
    c.alignment = Alignment(wrap_text=True, vertical="top")
for x in rows:
    ws.append([x["sku"], x["asin"], x["ean"], x["ean_src"]]
              + [x["imgs"].get(c, "") for c in img_cols] + [stamp_human])
# коды и штрихкоды — текстом: иначе Excel съест ведущий ноль у 08455000
for row in ws.iter_rows(min_row=2, max_col=3):
    for c in row:
        c.number_format = "@"
widths = [16, 13, 15, 16] + [40] * len(img_cols) + [16]
for i, w_ in enumerate(widths, start=1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w_
ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions

n = len(rows)
n_no_ean_amz = sum(1 for x in rows if x["ean_src"] != SRC_AMAZON)
n_ean_dir = sum(1 for x in rows if x["ean_src"] == SRC_DIR)
n_no_ean = sum(1 for x in rows if x["ean_src"] == SRC_NONE)
n_lt7 = sum(1 for x in rows if x["n_img"] < 7)
n_not_found = sum(1 for x in rows if not x["in_catalog"])
n_multi = sum(1 for x in rows if x["ean_multi"])
summary = [
    ("Export date", stamp_human + " (Kyiv time)"),
    ("Listings report date", listing_date),
    ("Products in export (active Amazon ES SKUs)", n),
    ("Unique ASINs", len({x["asin"] for x in rows if x["asin"]})),
    ("Excluded: amzn.gr.* (Amazon's resale of returned items, not our SKUs)", n_grade),
    ("ASIN not found in Amazon ES catalog", n_not_found),
    ("No EAN in Amazon catalog", n_no_ean_amz),
    ("  of which EAN taken from SKU directory", n_ean_dir),
    ("  of which no EAN anywhere", n_no_ean),
    ("Several EANs in Amazon catalog (comma-separated)", n_multi),
    ("Fewer than 7 images", n_lt7),
    ("Products by number of images (count: products)",
     ", ".join(f"{k}: {v}" for k, v in sorted({k: sum(1 for x in rows if x["n_img"] == k)
                                                for k in {x["n_img"] for x in rows}}.items()))),
    ("Highest additional image slot", f"PT{pt_max:02d}"),
    ("Other image variants (not in columns, e.g. swatches)",
     ", ".join(f"{k}: {v}" for k, v in sorted(other_variants.items())) or "none"),
    ("EAN source values", f"{SRC_AMAZON} — from Amazon catalog; {SRC_DIR} — from our SKU directory "
                          f"when Amazon has none; {SRC_NONE} — not found in either"),
    ("Images", "Main image first, then PT01, PT02… in Amazon's order; each link is the largest size Amazon returns"),
    ("Source", "SP-API Catalog Items 2022-04-01, Amazon.es, fresh request at export time"),
]
s2 = wb.create_sheet("Summary")
for k, v in summary:
    s2.append([k, v])
s2.column_dimensions["A"].width = 70
s2.column_dimensions["B"].width = 60
for c in s2["A"]:
    c.font = Font(bold=True)

import os
os.makedirs(OUT_DIR, exist_ok=True)
tmp = f"/tmp/amazon_es_catalog_{stamp}.xlsx"
wb.save(tmp)
name = f"amazon_es_catalog_{stamp}{'_probe' if LIMIT else ''}.xlsx"
shutil.copy(tmp, f"{OUT_DIR}/{name}")
if not LIMIT:
    shutil.copy(tmp, f"{OUT_DIR}/amazon_es_catalog_latest.xlsx")
print(f"✅ {OUT_DIR}/{name}")

result = {"file": f"{OUT_DIR}/{name}", "rows": n, "asins": len(asins), "not_found": n_not_found,
          "no_ean_amazon": n_no_ean_amz, "ean_from_dir": n_ean_dir, "no_ean": n_no_ean,
          "multi_ean": n_multi, "lt7": n_lt7, "pt_max": pt_max, "other_variants": other_variants,
          "listing_date": listing_date, "grade_excluded": int(n_grade), "no_asin_listings": len(no_asin),
          "img_hist": {k: sum(1 for x in rows if x["n_img"] == k) for k in sorted({x["n_img"] for x in rows})},
          "no_ean_list": [x["sku"] for x in rows if x["ean_src"] == SRC_NONE][:60]}
print(json.dumps(result, ensure_ascii=False))
dbutils.notebook.exit(json.dumps(result, ensure_ascii=False))  # noqa: F821
