# dict_kit.py — общие детали справочников: инструкция, шаблон Excel, загрузка с проверкой до сохранения
"""Один приём на все справочники (решение владельца 06.10.2026): наверху — короткая инструкция, таблица с
поиском и фильтрами, у строки — изменить / удалить (удаление с проверкой связей, иначе архив), под таблицей —
форма добавления с подсказками к каждому полю и массовая загрузка из Excel: шаблон, загрузка, проверка ДО
сохранения.

Здесь живут только общие части; что именно проверять и как записывать — дело раздела: он передаёт описание
полей и функцию проверки строки. Тексты приходят функцией `tr` раздела, чтобы язык жил в одном словаре.

Правила загрузки — те же, что у загрузки прогноза из файла (ТЗ 010 §11), чтобы человек не учил два набора:
пустая ячейка ничего не меняет и не обнуляет; ошибка хотя бы в одной строке блокирует сохранение целиком
(половина загруженного файла хуже незагруженного); перед сохранением видно «было → станет».
"""
from dataclasses import dataclass, field
from datetime import date, datetime
import io
import re

import pandas as pd
import streamlit as st

from util import grid_height


@dataclass
class Field:
    code: str                 # имя колонки в базе / ключ строки
    label: str                # заголовок в шаблоне и на экране
    hint: str                 # что вводить и откуда брать — и в форме, и в листе «Как заполнять»
    kind: str = "text"        # text | int | float | date | choice | ean
    required: bool = False
    choices: dict = field(default_factory=dict)   # код → подпись (для choice)
    example: str = ""


def instruction(md: str, caption: str = "") -> None:
    """Короткая инструкция наверху раздела: что это за справочник и для чего, и как с ним работать."""
    with st.container(border=True):
        st.markdown(md)
        if caption:
            st.caption(caption)


def template_bytes(fields: list, sheet: str, howto_title: str, cols_titles: tuple, rows: list | None = None) -> bytes:
    """Шаблон Excel: лист с заголовками (и, если передано, текущими строками) плюс лист «Как заполнять» —
    поле, обязательно ли, что вводить и откуда брать, пример."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.append([f.label for f in fields])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in rows or []:
        ws.append([r.get(f.code, "") for f in fields])
    for i, f in enumerate(fields, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(40, len(f.label) + 6))
    # текстовые колонки — текстом, иначе Excel съест ведущий ноль у артикула и EAN
    for i, f in enumerate(fields, start=1):
        if f.kind in ("text", "ean", "choice"):
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i, max_row=max(2, ws.max_row) + 500):
                for c in row:
                    c.number_format = "@"
    hw = wb.create_sheet(howto_title[:31])
    hw.append(list(cols_titles))
    for c in hw[1]:
        c.font = Font(bold=True)
    for f in fields:
        hw.append([f.label, "✓" if f.required else "", f.hint,
                   ", ".join(f.choices.values()) if f.choices else f.example])
    for col, w in zip("ABCD", (24, 12, 80, 30)):
        hw.column_dimensions[col].width = w
    for row in hw.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_upload(uploaded, fields: list) -> tuple:
    """Файл → таблица с кодами полей. Колонку узнаём по подписи (в любом регистре) или по коду поля.
    Возвращает (таблица, неизвестные колонки, отсутствующие обязательные)."""
    name = getattr(uploaded, "name", "").lower()
    if name.endswith(".csv"):
        raw = pd.read_csv(uploaded, dtype=str, keep_default_na=False)
    else:
        raw = pd.read_excel(uploaded, dtype=str, keep_default_na=False, sheet_name=0)
    norm = lambda s_: re.sub(r"\s+", " ", str(s_ or "")).strip().lower()
    by = {}
    for f in fields:
        by[norm(f.label)] = f.code
        by[norm(f.code)] = f.code
    mapping, unknown = {}, []
    for c in raw.columns:
        code = by.get(norm(c))
        if code:
            mapping[c] = code
        elif norm(c) and not norm(c).startswith("unnamed"):
            unknown.append(str(c))
    df = raw.rename(columns=mapping)[list(dict.fromkeys(mapping.values()))]
    df = df.apply(lambda col: col.map(lambda v: str(v).strip() if v is not None else ""))
    df = df[(df != "").any(axis=1)].reset_index(drop=True)
    missing = [f.label for f in fields if f.required and f.code not in df.columns]
    return df, unknown, missing


def parse_value(f: Field, v: str):
    """Значение ячейки → (значение, ошибка). Пусто → (None, None): пустая ячейка ничего не меняет."""
    s_ = str(v or "").strip()
    if not s_:
        return None, None
    if f.kind == "int":
        try:
            x = float(s_.replace(",", ".").replace(" ", ""))
            return (int(round(x)), None) if x >= 0 else (None, "neg")
        except ValueError:
            return None, "num"
    if f.kind == "float":
        try:
            x = float(s_.replace(",", ".").replace(" ", ""))
            return (x, None) if x >= 0 else (None, "neg")
        except ValueError:
            return None, "num"
    if f.kind == "date":
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"):
            try:
                return datetime.strptime(s_, fmt).date(), None
            except ValueError:
                continue
        return None, "date"
    if f.kind == "ean":
        d = re.sub(r"\.0$", "", s_)
        return (d, None) if re.fullmatch(r"\d{8}|\d{12,14}", d) else (None, "ean")
    if f.kind == "choice":
        low = s_.lower()
        for code, lbl in f.choices.items():
            if low in (code.lower(), str(lbl).lower()):
                return code, None
        return None, "choice"
    return s_, None


def check_view(errors: list, changes: list, n_new: int, n_upd: int, n_same: int, tr) -> bool:
    """Показ результата проверки. Возвращает True, если сохранять можно (ошибок нет и есть что сохранять)."""
    c1, c2, c3, c4 = st.columns(4)
    c1.metric(tr("kit_new"), n_new)
    c2.metric(tr("kit_upd"), n_upd)
    c3.metric(tr("kit_same"), n_same)
    c4.metric(tr("kit_err"), len(errors))
    if errors:
        st.error(tr("kit_err_block"))
        st.dataframe(pd.DataFrame(errors, columns=[tr("kit_c_row"), tr("kit_c_key"), tr("kit_c_field"),
                                                   tr("kit_c_value"), tr("kit_c_why")]),
                     hide_index=True, use_container_width=True)
        return False
    if changes:
        st.dataframe(pd.DataFrame(changes, columns=[tr("kit_c_key"), tr("kit_c_field"), tr("kit_c_old"), tr("kit_c_new")]),
                     hide_index=True, use_container_width=True, height=min(420, 38 + 35 * len(changes)))
    if not (n_new or n_upd):
        st.info(tr("kit_nothing"))
        return False
    return True


# ── список → карточка: общий приём для всех справочников (замечания веб-агента 06.10.2026) ──────────────

def search_input(label: str, key: str, tr, example: str = "", where=None) -> str:
    """Поле поиска по справочнику. Streamlit применяет введённый текст по Enter или когда фокус уходит из
    поля — «по мере ввода» он не умеет, поэтому это сказано прямо в поле, а не оставлено догадке."""
    ph = (example + " · " if example else "") + tr("kit_search_enter")
    # `where` — колонка или контейнер, куда ставить поле: без него оно уходит в общий поток страницы
    return (where or st).text_input(label, key=key, placeholder=ph, help=tr("kit_search_help")).strip()


def table_pick(df: pd.DataFrame, ids: list, key: str, fsig, column_config: dict, height: int | None = None):
    """Список справочника; строка выбирается кликом по любой её ячейке.

    Режим «single-cell» без «single-row»: строчный режим рисует слева колонку флажков, а массовых действий у
    списков нет — флажки ничего не делали и только сбивали с толку. Ключ таблицы зависит от фильтров
    (`fsig`): выбор хранится номером строки, и после смены фильтра тот же номер указал бы на чужую запись.

    Возвращает id выбранной строки, только если выбор в таблице ИЗМЕНИЛСЯ в этом прогоне, иначе None —
    иначе однажды выбранная строка перебивала бы каждый следующий выбор в поле карточки."""
    fkey = abs(hash(fsig)) % 10 ** 8
    ev = st.dataframe(df, key=f"{key}_{fkey}", use_container_width=True, hide_index=True,
                      height=height or grid_height(len(df), 11),
                      on_select="rerun", selection_mode="single-cell", column_config=column_config)
    sel = ev.selection if ev is not None and hasattr(ev, "selection") else {}
    cells = list(sel.get("cells", []) or [])
    pos = int(cells[0][0]) if cells else None
    sig = (fkey, tuple(tuple(c) for c in cells))
    picked = None
    if pos is not None and pos < len(ids) and sig != st.session_state.get(f"{key}__sig"):
        picked = ids[pos]
    st.session_state[f"{key}__sig"] = sig
    return picked


def card_focus(state_key: str, item_id) -> None:
    """Открыть карточку записи (после создания, после клика в таблице). Можно звать в любой момент
    прогона: пишет в свой ключ состояния, а не в ключ виджета."""
    st.session_state[state_key] = item_id
    st.session_state[f"{state_key}__gen"] = st.session_state.get(f"{state_key}__gen", 0) + 1


def card_pick(label: str, ids: list, labels: dict, state_key: str, picked=None):
    """Поле «какая карточка открыта».

    Открытая запись живёт в своём ключе (`state_key`), а не в ключе виджета. Раньше поле было виджетом
    с постоянным ключом, и при смене фильтров (список вариантов другой) Streamlit пересоздавал его с первым
    вариантом — карточка сбрасывалась, хотя её строка оставалась в таблице. Теперь открытая карточка
    остаётся, пока её строка видна; ушла из таблицы — открывается первая строка."""
    if picked is not None:
        card_focus(state_key, picked)
    cur = st.session_state.get(state_key)
    if cur not in ids:
        cur = ids[0]
    gen = st.session_state.get(f"{state_key}__gen", 0)
    wkey = f"{state_key}__box_{abs(hash(tuple(ids))) % 10 ** 8}_{gen}"
    chosen = st.selectbox(label, ids, index=ids.index(cur), key=wkey,
                          format_func=lambda i: labels.get(i, str(i)))
    st.session_state[state_key] = chosen
    return chosen
